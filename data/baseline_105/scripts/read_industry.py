# -*- coding: utf-8 -*-
import sys, os, json, glob
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from collections import Counter

buf = []
def P(*a): buf.append(' '.join(str(x) for x in a))

# 1) 行业分类清单.xlsx
XL = r'[WORKSPACE]/行业分类清单.xlsx'
P('=' * 70); P('行业分类清单.xlsx'); P('=' * 70)
try:
    import openpyxl
    wb = openpyxl.load_workbook(XL, data_only=True)
    for ws in wb.worksheets:
        P('--- sheet: %s  (%dx%d) ---' % (ws.title, ws.max_row, ws.max_column))
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 110:
                P('   ...(截断)')
                break
            vals = ['' if v is None else str(v).strip() for v in row]
            if any(vals):
                P('   ', ' | '.join(vals))
except Exception as e:
    P('ERR', type(e).__name__, e)

# 2) PDF 根目录结构
ROOT = r'[DATA_ROOT]/ESG报告'
P(''); P('=' * 70); P('ESG报告 目录结构'); P('=' * 70)
tot = 0
for d in sorted(os.listdir(ROOT)):
    p = os.path.join(ROOT, d)
    if os.path.isdir(p):
        n = len(glob.glob(os.path.join(p, '*.pdf')))
        tot += n
        P('  %-24s %d 个 pdf' % (d, n))
    else:
        P('  [file] %s' % d)
P('  合计 pdf =', tot)

open(r'[REPO]/data/baseline_105/scripts/_industry.txt',
     'w', encoding='utf-8').write('\n'.join(buf))
print('OK', len(buf), 'lines')
